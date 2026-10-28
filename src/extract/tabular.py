"""CSV / Excel metadata via pandas."""
from __future__ import annotations

import io
import re
import warnings
from pathlib import Path

import pandas as pd

from . import FileMetadata

SAMPLE_ROWS = 5000
DATE_NAME_RE = re.compile(r"date|time|day|month|period|timestamp|acquired|observed", re.I)
YEAR_NAME_RE = re.compile(r"^(year|yr|yyyy)$", re.I)
LAT_WORDS = {"lat", "latitude"}
LON_WORDS = {"lon", "long", "longitude", "lng"}
BLANK_RE = re.compile(r"^[\s,;]*$")


def _read_text(path: Path) -> str:
    raw = path.read_bytes()[: 32 * 1024 * 1024]
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
    return "\n".join(text.splitlines()[:SAMPLE_ROWS])


def _read_csv(path: Path) -> list[pd.DataFrame]:
    """CSV sample as one frame per block (some exports stack several tables split by blank lines)."""
    blocks: list[list[str]] = []
    current: list[str] = []
    for line in _read_text(path).split("\n"):
        if BLANK_RE.match(line):
            if current:
                blocks.append(current)
            current = []
        else:
            current.append(line)
    if current:
        blocks.append(current)
    frames = []
    for block in blocks:
        text = "\n".join(block)
        try:
            frames.append(pd.read_csv(io.StringIO(text), sep=None, engine="python"))
        except Exception:
            frames.append(pd.read_csv(io.StringIO(text), sep=None, engine="python", on_bad_lines="skip"))
    return frames


def _count_rows(path: Path) -> int:
    count = 0
    with open(path, "rb") as fh:
        while chunk := fh.read(1 << 20):
            count += chunk.count(b"\n")
    return max(0, count - 1)


def _is_text(series: pd.Series) -> bool:
    return pd.api.types.is_object_dtype(series) or pd.api.types.is_string_dtype(series)


def _words(name) -> set[str]:
    return set(re.findall(r"[a-z]+", str(name).lower()))


def _kv_pairs(df: pd.DataFrame) -> dict[str, object] | None:
    """A 'key,value[,unit]' table -> {key: value}; None for regular tables."""
    if df.shape[1] < 2 or df.shape[1] > 3 or len(df) > 500:
        return None
    keys = df.iloc[:, 0]
    if not _is_text(keys) or keys.isna().any() or keys.nunique() != len(keys):
        return None
    return dict(zip(keys.astype(str), df.iloc[:, 1]))


def _date_range(df: pd.DataFrame) -> tuple[pd.Timestamp | None, pd.Timestamp | None]:
    lo = hi = None
    for col in df.columns:
        series = df[col].dropna()
        if series.empty:
            continue
        parsed = None
        if pd.api.types.is_datetime64_any_dtype(series):
            parsed = series
        elif YEAR_NAME_RE.match(str(col)) and pd.api.types.is_numeric_dtype(series):
            years = series[(series >= 1800) & (series <= 2100)]
            if len(years) >= 0.8 * len(series):
                parsed = pd.to_datetime(years.astype(int).astype(str), format="%Y", errors="coerce")
        elif DATE_NAME_RE.search(str(col)) and _is_text(series):
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                candidate = pd.to_datetime(series.astype(str), errors="coerce", utc=True)
            if candidate.notna().mean() >= 0.8:
                parsed = candidate.dt.tz_localize(None)
        if parsed is None:
            continue
        parsed = parsed.dropna()
        if parsed.empty:
            continue
        lo = parsed.min() if lo is None else min(lo, parsed.min())
        hi = parsed.max() if hi is None else max(hi, parsed.max())
    return lo, hi


def _bbox(df: pd.DataFrame, kv: dict | None):
    if kv is not None:
        lat = next((v for k, v in kv.items() if _words(k) & LAT_WORDS), None)
        lon = next((v for k, v in kv.items() if _words(k) & LON_WORDS), None)
        la = pd.to_numeric(pd.Series([lat]), errors="coerce")
        lo = pd.to_numeric(pd.Series([lon]), errors="coerce")
    else:
        lat = next((c for c in df.columns if _words(c) & LAT_WORDS), None)
        lon = next((c for c in df.columns if _words(c) & LON_WORDS), None)
        if lat is None or lon is None:
            return None
        la, lo = pd.to_numeric(df[lat], errors="coerce"), pd.to_numeric(df[lon], errors="coerce")
    la, lo = la.dropna(), lo.dropna()
    if la.empty or lo.empty or la.abs().max() > 90 or lo.abs().max() > 180:
        return None
    return (float(lo.min()), float(la.min()), float(lo.max()), float(la.max()))


def extract_tabular(path: Path) -> FileMetadata:
    ext = path.suffix.lower()
    frames: list[pd.DataFrame] = []
    extra: dict = {}
    if ext == ".csv":
        frames.extend(_read_csv(path))
        extra["row_count"] = _count_rows(path)
    else:
        sheets = pd.read_excel(path, sheet_name=None, nrows=SAMPLE_ROWS)
        frames.extend(sheets.values())
        extra["sheet_names"] = list(sheets.keys())
        if ext in (".xlsx", ".xlsm"):
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            try:
                extra["row_count"] = sum(max(0, (ws.max_row or 1) - 1) for ws in wb.worksheets)
            finally:
                wb.close()
        else:
            extra["row_count"] = sum(len(f) for f in frames)

    fields: list[str] = []
    lo = hi = None
    bbox = None
    for df in frames:
        kv = _kv_pairs(df)
        names = list(kv.keys()) if kv is not None else [str(c) for c in df.columns]
        fields.extend(n for n in names if n not in fields)
        a, b = _date_range(df)
        if a is not None:
            lo = a if lo is None else min(lo, a)
            hi = b if hi is None else max(hi, b)
        bbox = bbox or _bbox(df, kv)

    return FileMetadata(
        kind="tabular",
        fields=fields,
        time_start=lo.date().isoformat() if lo is not None else None,
        time_end=hi.date().isoformat() if hi is not None else None,
        bbox=bbox,
        crs="EPSG:4326" if bbox else None,
        extra=extra,
    )
